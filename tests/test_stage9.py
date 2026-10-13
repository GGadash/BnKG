import csv
import json
import threading
from dataclasses import replace

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QFileDialog
from test_history import create_log

from bnkg.core.export import export_session
from bnkg.core.export_profiles import (
    ALL_IDS,
    BUILTINS,
    KEY,
    decode_profiles,
    describe,
    encode_profiles,
    profile,
)
from bnkg.ui.export_profiles import ExportPresets
from bnkg.ui.window import MainWindow


@pytest.mark.parametrize("fmt", ["csv", "tsv", "txt", "json", "ndjson", "xlsx", "html"])
def test_preview_matches_format_payload(tmp_path, fmt):
    source = tmp_path / "source.sqlite"
    store, sample = create_log(source, 0, close=False)
    store.append(replace(sample, values={"LAF": 0, "LAeq": -5, "EXTRA": 23}))
    store.event("fixture", "local event")
    store.close()
    before = source.read_bytes()
    config = profile(fmt, ["LAF"], True, True)
    text = describe(config)
    target = tmp_path / ("out." + fmt)
    assert export_session(source, target, fmt, **config.options()) == 1
    assert source.read_bytes() == before
    if fmt in {"csv", "tsv", "txt"}:
        with target.open(encoding="utf-8", newline="") as handle:
            rows = list(csv.reader(handle, delimiter="," if fmt == "csv" else "\t"))
        assert rows[0] == ["Timestamp_UTC", "Elapsed_s", "State", "Origin", "Quality", "LAF_dB"]
        assert not config.include_raw and not config.include_spectra
        assert "Scalar table only" in text
    elif fmt in {"json", "ndjson"}:
        data = (
            json.loads(target.read_text())
            if fmt == "json"
            else [json.loads(line) for line in target.read_text().splitlines()]
        )
        row = data["samples"][0] if fmt == "json" else data[1]
        assert row["values"] == {"LAF": 0} and "raw_values" in row and "spectra" in row
        assert "independently of scalar selection" in text
        assert "Session event records are not included" in text
    elif fmt == "xlsx":
        wb = load_workbook(target, read_only=True)
        try:
            assert {"Summary", "Parameters", "Time Series", "Spectra", "Events"} == set(
                wb.sheetnames
            )
            assert "Includes the Events sheet" in text
        finally:
            wb.close()
    else:
        html = target.read_text(encoding="utf-8")
        assert "Session and parameter context" in html and "Direct saved spectra" in html
        assert "Session event records are not included" in text
    assert "integration intervals" in text and "Options preview only" in text


def test_profile_storage_bounds_semantics_and_raw_independence():
    saved = {"My JSON": profile("json", ALL_IDS, True, True)}
    assert decode_profiles(encode_profiles(saved)) == saved
    assert "additional stored scalar keys" in describe(saved["My JSON"])
    assert "formula details are not embedded" in describe(profile("csv", ["LCeq_minus_LAeq"]))
    for options in (
        ("pdf", ["LAF"], False, False),
        ("csv", [], False, False),
        ("json", ["unknown"], False, False),
        ("json", ["LAF"], 1, False),
    ):
        with pytest.raises(ValueError):
            profile(*options)
    with pytest.raises(ValueError):
        encode_profiles({f"Name {i}": saved["My JSON"] for i in range(21)})
    with pytest.raises(ValueError):
        decode_profiles('{"version":true,"presets":[]}')
    payload = json.loads(encode_profiles(saved))
    payload["presets"][0]["profile"]["format_name"] = "csv"
    with pytest.raises(ValueError, match="unsupported by its format"):
        decode_profiles(json.dumps(payload))


def test_preset_preview_persistence_and_invalid_storage(qtbot, tmp_path):
    prefs = QSettings(str(tmp_path / "prefs.ini"), QSettings.Format.IniFormat)
    current = profile("ndjson", ["LAF", "LAeq"], True, False)
    widget = ExportPresets(prefs, lambda: current)
    qtbot.addWidget(widget)
    applied = []
    widget.apply_requested.connect(applied.append)
    widget.choices.setCurrentText("Measurement report")
    assert not applied
    widget.apply.click()
    assert applied == [BUILTINS["Measurement report"]]
    widget.name.setText("Saved field options")
    widget.save.click()
    assert decode_profiles(prefs.value(KEY)) == {"Saved field options": current}
    widget.name.setText("saved FIELD options")
    widget.save.click()
    assert "already exists" in widget.status.text()
    restored = ExportPresets(prefs, lambda: current)
    qtbot.addWidget(restored)
    assert restored.saved == widget.saved
    raw = '{"version":9,"presets":[]}'
    prefs.setValue(KEY, raw)
    broken = ExportPresets(prefs, lambda: current)
    qtbot.addWidget(broken)
    assert broken.error and not broken.save.isEnabled() and prefs.value(KEY) == raw


def test_export_job_captures_options_and_restores_controls(qtbot, tmp_path, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    captured = {}

    def slow(database, target, format_name, **options):
        captured.update(database=database, format_name=format_name, **options)
        entered.set()
        release.wait(3)
        return 1

    monkeypatch.setattr("bnkg.ui.exporting.export_session", slow)
    window = MainWindow(data_dir=tmp_path / "app")
    commands = []
    window.command_requested.connect(lambda *args: commands.append(args))
    try:
        wanted = profile("json", ["LAF"], True, False)
        window.apply_export_profile(wanted)
        window.start_export(tmp_path / "original.sqlite", tmp_path / "out.json", "json")
        qtbot.waitUntil(entered.is_set)
        window.apply_export_profile(BUILTINS["Scalar table"])
        window.export_parameters.append(
            "LAeq"
        )  # External mutation must not affect captured job options.
        window.receive_session(str(tmp_path / "next.sqlite"), True, 2)
        assert not window.format_box.isEnabled() and not window.export_presets.isEnabled()
        assert "original.sqlite" in window.export_preview.toPlainText()
        assert "next.sqlite" not in window.export_preview.toPlainText()
        assert (
            captured["parameter_ids"] == ["LAF"]
            and captured["include_spectra"]
            and not captured["include_raw"]
        )
        assert not commands
        release.set()
        qtbot.waitUntil(lambda: window.export_job is None)
        assert window.format_box.isEnabled() and window.export_presets.isEnabled()
        assert "next.sqlite" in window.export_preview.toPlainText()
        window.export_parameters = []
        errors = []
        monkeypatch.setattr(window, "show_error", errors.append)
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            lambda *args: pytest.fail("No dialog for invalid selection"),
        )
        window.export_current()
        assert errors and window.export_job is None
    finally:
        release.set()
        window.close()
        qtbot.waitUntil(lambda: window.thread_finished)

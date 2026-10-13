import json
import math
import sqlite3
import threading
from dataclasses import replace
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QMetaObject, Qt

from bnkg.core.device import DemoDevice
from bnkg.core.export import ExportCancelled, export_session
from bnkg.core.sessions import SessionStore
from bnkg.core.trends import TrendPoint, series_points
from bnkg.core.webxi import WebXiClient
from bnkg.ui.window import MainWindow


def log(path, count=3):
    store = SessionStore(path, "instrument")
    sample = replace(
        DemoDevice(seed=1).read(), origin="instrument", quality="synthetic test fixture"
    )
    for _ in range(count):
        store.append(sample)
    return store, sample


def test_export_releases_source_before_formatting_and_uses_snapshot(tmp_path):
    path, target = tmp_path / "source.sqlite", tmp_path / "out.json"
    store, sample = log(path, 101)
    appended = []

    def progress(done, total, phase):
        if phase.startswith("Writing") and not appended:
            store.append(sample)  # Would lock/fail if formatting held a source read transaction.
            appended.append(True)

    try:
        assert export_session(path, target, "json", progress=progress) == 101
        assert appended and store.count == 102
        assert len(json.loads(target.read_text())["samples"]) == 101
    finally:
        store.close()
    assert not list(tmp_path.glob(".bnkg-export-*"))


@pytest.mark.parametrize("format_name", ["csv", "xlsx", "json", "html"])
def test_cancel_removes_partial_and_snapshot_without_publishing(tmp_path, format_name):
    path, target = tmp_path / "source.sqlite", tmp_path / ("out." + format_name)
    store, _ = log(path, 101)
    store.close()
    cancel = False

    def progress(done, total, phase):
        nonlocal cancel
        if phase.startswith("Writing"):
            cancel = True

    with pytest.raises(ExportCancelled):
        export_session(path, target, format_name, cancelled=lambda: cancel, progress=progress)
    assert not target.exists()
    assert not list(tmp_path.glob("*.partial")) and not list(tmp_path.glob(".bnkg-export-*"))


def test_selected_parameters_and_html_escaping(tmp_path):
    path = tmp_path / "source.sqlite"
    store, sample = log(path, 0)
    store.append(
        replace(sample, state='<script>alert("x")</script>', values={"LAF": 61, "LAeq": None})
    )
    store.close()
    target = tmp_path / "out.json"
    export_session(path, target, "json", parameter_ids=["LAF"], include_raw=False)
    data = json.loads(target.read_text())
    assert data["samples"][0]["values"] == {"LAF": 61}
    assert "raw_values" not in data["samples"][0]
    html = tmp_path / "report.html"
    export_session(path, html, "html")
    text = html.read_text(encoding="utf-8")
    assert "<script>" not in text and "&lt;script&gt;" in text
    assert "Unavailable" in text and "Calculated" in text


def test_spectral_exports_preserve_scope_axes_nulls_and_context(tmp_path):
    path = tmp_path / "source.sqlite"
    descriptor = {
        "Name": "CPBLZF",
        "LocalName": "LZF",
        "FunctionType": "OctaveL",
        "AveragingMode": "Fast",
        "AcousticalWeighting": "Z",
        "DataType": "Int16",
        "Scale": 0.01,
        "Unit": "dB re 20uPa",
        "PeriodTime": 1073741824,
        "DataAxisType": "CPB",
        "DataAxisCpbBaseSystem": 10,
        "VectorLength": 3,
        "DataAxisCpbFirstBand": 28,
        "DataAxisCpbNumberFractions": 3,
    }
    raw = {
        "CPBLZF250ms": [0, -500, -32768],
        "_hbkg_context": {
            "sequences": {"CPBLZF": descriptor},
            "setup_basis": {},
            "metadata": {"MeasurementCounter": {"value": 9}},
        },
    }
    store, sample = log(path, 0)
    store.append(replace(sample, raw_values=raw))
    store.close()
    target = tmp_path / "bands.json"
    export_session(path, target, "json", include_spectra=True, include_raw=False)
    spectra = json.loads(target.read_text())["samples"][0]["spectra"]
    assert spectra["measurement_context"]["MeasurementCounter"]["value"] == 9
    series = spectra["series"][0]
    assert series["values"] == [0, -5, None] and series["meaning"]["kind"] == "current"
    assert series["meaning"]["period_ms"] is None and series["meaning"]["output_ms"] == 250
    assert series["descriptor"] == descriptor and len(spectra["series"]) == 1
    book = tmp_path / "bands.xlsx"
    export_session(path, book, "xlsx", include_spectra=True)
    with_book = load_workbook(book, read_only=True)
    assert "Spectra" in with_book.sheetnames
    rows = list(with_book["Spectra"].values)
    assert len(rows) == 4 and rows[1][10] == 0 and rows[3][10] is None
    with_book.close()


def test_trend_gaps_nulls_origins_and_unknown_timezone():
    start = datetime(2026, 9, 28, tzinfo=UTC)
    points = [
        TrendPoint.create(start + timedelta(seconds=t), {"LAF": v}, origin, "Running", "valid")
        for t, v, origin in [
            (0, 0, "instrument"),
            (1, None, "instrument"),
            (2, -5, "instrument"),
            (10, 60, "instrument"),
            (11, 61, "simulation"),
        ]
    ]
    x, y = series_points(points, "LAF")
    assert len(x) == 7 and sum(math.isnan(v) for v in y) == 3
    assert 0 in y and -5 in y
    assert TrendPoint.create("2026-09-28T12:00:00", {}, "instrument", "", "valid") is None


def test_capture_clear_and_bounds_do_not_clear_saved_log(qtbot, tmp_path):
    window = MainWindow(data_dir=tmp_path / "app")
    path = tmp_path / "log.sqlite"
    store, sample = log(path, 6)
    store.close()
    try:
        window.record_limit.setValue(3)
        window.auto_scroll.setChecked(False)
        for i in range(6):
            window.receive_captured(
                replace(sample, timestamp=sample.timestamp + timedelta(seconds=i))
            )
        assert window.table.rowCount() == 3 and len(window.capture_chart.points) == 3
        window.clear_capture_view()
        assert (
            window.table.rowCount() == 0 and not window.capture_chart.points and not window.history
        )
        assert window.capture_chart.points.maxlen == 3
        window.history.extend([(0, 0), (1, float("nan")), (8, -5)])
        window.draw_capture_trend()
        x, y = window.curve.getData()
        assert list(x) == [0, 1, 8, 8]
        assert y[0] == 0 and y[-1] == -5 and sum(math.isnan(v) for v in y) == 2
        with sqlite3.connect(path) as db:
            assert db.execute("SELECT count(*) FROM samples").fetchone()[0] == 6
    finally:
        window.close()
        qtbot.waitUntil(lambda: window.thread_finished)


def test_export_worker_does_not_block_acquisition_and_close_waits(qtbot, tmp_path, monkeypatch):
    import bnkg.ui.worker as worker_module

    entered, release = threading.Event(), threading.Event()

    class FixtureMeter(WebXiClient):
        def __init__(self, *args, **kwargs):
            super().__init__("fixture.local")
            self.demo = DemoDevice(seed=1)

        def identify(self):
            return {"HostName": "Offline fixture"}

        def read(self):
            return replace(self.demo.read(), origin="instrument", quality="synthetic test fixture")

    def slow_export(*args, **kwargs):
        entered.set()
        release.wait(3)
        if kwargs["cancelled"]():
            raise ExportCancelled()
        return 1

    monkeypatch.setattr(worker_module, "WebXiClient", FixtureMeter)
    monkeypatch.setattr("bnkg.ui.exporting.export_session", slow_export)
    window = MainWindow(data_dir=tmp_path)
    try:
        window.select_source("connect", {"address": "fixture.local"})
        qtbot.waitUntil(lambda: window.latest is not None)
        window.toggle_logging()
        qtbot.waitUntil(lambda: window.logging)
        window.start_export(window.session_path, tmp_path / "out.json", "json")
        qtbot.waitUntil(entered.is_set)
        for _ in range(3):
            QMetaObject.invokeMethod(window.worker, "poll", Qt.ConnectionType.QueuedConnection)
        qtbot.waitUntil(lambda: window.table.rowCount() >= 3)
        assert window.export_job is not None and window.logging
        window.close()
        assert not window.thread_finished
        release.set()
        qtbot.waitUntil(lambda: window.thread_finished)
        assert window.export_job is None
    finally:
        release.set()
        window.close()
        qtbot.waitUntil(lambda: window.thread_finished)


def test_large_xlsx_context_round_trips_without_truncation(tmp_path, monkeypatch):
    path = tmp_path / "source.sqlite"
    store, _ = log(path, 1)
    store.close()
    original = path.read_bytes()
    large = {"private_note": "=do not evaluate\n" + "α😀" * 20000}
    payload = {
        "measurement_context": large,
        "series": [
            {
                "key": "fixture",
                "meaning": {
                    "quantity": "LZF",
                    "kind": "current",
                    "weighting": "Z",
                    "period_ms": None,
                    "output_ms": 250,
                },
                "frequencies": [100, 125],
                "values": [0, None],
                "broadband": None,
                "descriptor": {"note": "x" * 35000},
            }
        ],
    }
    monkeypatch.setattr("bnkg.core.export.spectral_payload", lambda *a: payload)
    target = tmp_path / "large.xlsx"
    export_session(path, target, "xlsx", include_spectra=True)
    book = load_workbook(target, read_only=True)
    try:
        bands = list(book["Spectra"].values)
        chunks = list(book["Text chunks"].values)[1:]
        for column, expected in ((12, payload["series"][0]["descriptor"]), (13, large)):
            reference = bands[1][column].split(":", 1)[1].split(" ", 1)[0]
            parts = [r for r in chunks if r[0] == reference]
            assert [r[1] for r in parts] == list(range(1, len(parts) + 1))
            assert all(
                r[2] == len(parts) and len(r[3].encode("utf-16-le")) // 2 < 32000 for r in parts
            )
            assert json.loads("".join(json.loads(r[3]) for r in parts)) == expected
            assert bands[2][column] == bands[1][column]
        assert bands[1][10] == 0 and bands[2][10] is None
        assert path.read_bytes() == original
    finally:
        book.close()

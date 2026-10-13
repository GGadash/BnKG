import csv
import json
import sqlite3
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from bnkg.core.device import DemoDevice
from bnkg.core.export import export_session
from bnkg.core.parameters import Sample, decode_laf, differences
from bnkg.core.safety import SafetyLevel, classify, require_allowed
from bnkg.core.sessions import SessionStore
from bnkg.core.webxi import LAF, ROOT, WebXiClient, WebXiError, normalize_address


@pytest.mark.parametrize(
    "address,expected",
    [
        ("192.168.1.20", "http://192.168.1.20"),
        (" slm.local ", "http://slm.local"),
        ("https://slm.local:8443/", "https://slm.local:8443"),
        ("::1", "http://[::1]"),
    ],
)
def test_address(address, expected):
    assert normalize_address(address) == expected


@pytest.mark.parametrize(
    "address",
    [
        "",
        "http://host/webxi",
        "ftp://host",
        "host?Action=StartPause",
        "user:pass@host",
        "host\\path",
        "host:0",
        "999.1.2.3",
        "bad host",
        "-bad.local",
        "host#fragment",
    ],
)
def test_reject_unsafe_or_malformed_addresses(address):
    with pytest.raises(ValueError):
        normalize_address(address)


@pytest.mark.parametrize("raw", [None, True, "6400", {}, float("nan"), float("inf")])
def test_laf_rejects_ambiguous_values(raw):
    with pytest.raises(ValueError):
        decode_laf(raw)


def test_laf_scaling_and_missing_derivations():
    assert decode_laf(6425) == 64.25
    values = differences({"LAeq": 60.0, "LCeq": 67.2, "LAIeq": None, "LAFeq": 61.0})
    assert values["LCeq_minus_LAeq"] == pytest.approx(7.2)
    assert values["LAIeq_minus_LAFeq"] is None
    assert differences({})["LCeq_minus_LAeq"] is None


@pytest.mark.parametrize(
    "operation", ["calibration", "factory_reset", "set_ip", "unknown", "write", "key_menu"]
)
def test_no_real_device_writes_allowed(operation):
    assert classify(operation).level == SafetyLevel.BLOCKED
    with pytest.raises(PermissionError):
        require_allowed(operation)


def test_demo_pause_freezes_measurement_values_not_receipt_time():
    demo = DemoDevice()
    before = demo.read()
    demo.toggle()
    paused = demo.read()
    assert paused.origin == "simulation"
    assert paused.quality == "synthetic"
    assert paused.values == before.values
    assert paused.elapsed_s == before.elapsed_s
    assert paused.state == "Paused"
    demo.toggle()
    assert demo.read().elapsed_s > paused.elapsed_s


class FakeReply:
    status_code = 200

    def __init__(self, payload):
        self.payload = payload

    def __enter__(self):
        return self

    def __exit__(self, *_):
        pass

    def raise_for_status(self):
        pass

    def iter_content(self, chunk_size):
        yield self.payload


class FakeSession:
    def __init__(self, payload):
        self.payload = payload
        self.calls = []

    def get(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return FakeReply(self.payload)

    def close(self):
        pass


def test_client_read_contract_and_bounded_payload():
    session = FakeSession(b"6450")
    client = WebXiClient("slm.local", session=session)
    assert client.get(LAF) == 6450
    url, options = session.calls[0]
    assert url.endswith(LAF)
    assert options["allow_redirects"] is False
    assert options["timeout"] == (2.0, 3.0)
    assert not session.trust_env
    with pytest.raises(ValueError):
        client.get("/webxi?Action=StartPause")
    with pytest.raises(ValueError):
        client.get(ROOT, recursive=True)
    session.payload = b"x" * 2_000_001
    with pytest.raises(WebXiError, match="size/time"):
        client.get(ROOT)


@pytest.mark.parametrize("payload", [b"<html>Login</html>", b"\xff"])
def test_client_rejects_non_json(payload):
    with pytest.raises(WebXiError, match="JSON"):
        WebXiClient("slm.local", FakeSession(payload)).get(ROOT)


def make_session(path):
    store = SessionStore(path, "simulation")
    sample = Sample(
        datetime(2026, 9, 17, 12, 0, tzinfo=UTC),
        1.0,
        "=unsafe",
        "simulation",
        {"LAF": 64.25, "LAeq": None},
        {},
        "synthetic",
    )
    store.append(sample)
    store.event("test", "=not_a_formula")
    store.close()
    return sample


@pytest.mark.parametrize("format_name", ["csv", "tsv", "txt", "json", "ndjson", "xlsx"])
def test_exports_preserve_values_and_origin(tmp_path, format_name):
    database = tmp_path / "source.sqlite"
    make_session(database)
    target = tmp_path / ("export." + format_name)
    assert export_session(database, target, format_name) == 1
    if format_name in {"csv", "tsv", "txt"}:
        with target.open(encoding="utf-8", newline="") as handle:
            row = next(csv.DictReader(handle, delimiter="," if format_name == "csv" else "\t"))
        assert row["LAF_dB"] == "64.25"
        assert row["LAeq_dB"] == ""
        assert row["Origin"] == "simulation"
        assert row["State"] == "'=unsafe"
    elif format_name == "json":
        result = json.loads(target.read_text())
        assert result["samples"][0]["values"]["LAeq"] is None
        assert result["session"]["origin"] == "simulation"
        assert result["parameters"][9]["source"] == "simulation"
    elif format_name == "ndjson":
        lines = [json.loads(line) for line in target.read_text().splitlines()]
        assert lines[0]["type"] == "metadata"
        assert lines[1]["values"]["LAF"] == 64.25
    else:
        workbook = load_workbook(target, read_only=True)
        assert workbook.sheetnames == ["Summary", "Parameters", "Time Series", "Events"]
        assert workbook["Time Series"].cell(2, 3).value == "'=unsafe"
        assert workbook["Time Series"].cell(2, 4).value == "simulation"
        assert workbook["Events"].cell(2, 3).data_type != "f"
        workbook.close()
    with pytest.raises(FileExistsError):
        export_session(database, target, format_name)


def test_source_and_existing_database_are_preserved(tmp_path):
    database = tmp_path / "session.sqlite"
    make_session(database)
    original = database.read_bytes()
    with pytest.raises(FileExistsError):
        SessionStore(database, "simulation")
    with pytest.raises(ValueError):
        export_session(database, database, "json")
    assert database.read_bytes() == original
    with sqlite3.connect(database) as db:
        assert db.execute("SELECT count(*) FROM samples").fetchone()[0] == 1
        assert db.execute("SELECT value FROM metadata WHERE key='closed_utc'").fetchone()


def test_client_rejects_redirect_without_following(monkeypatch):
    monkeypatch.setattr(FakeReply, "status_code", 302)
    session = FakeSession(b"{}")
    with pytest.raises(WebXiError, match="redirect refused"):
        WebXiClient("slm.local", session).get(ROOT)
    assert len(session.calls) == 1


def test_client_timeout_is_reported_without_retry():
    import requests

    class TimeoutSession(FakeSession):
        def get(self, url, **kwargs):
            self.calls.append((url, kwargs))
            raise requests.Timeout("instrument not responding")

    session = TimeoutSession(b"")
    with pytest.raises(WebXiError, match="request timed out"):
        WebXiClient("slm.local", session).get(ROOT)
    assert len(session.calls) == 1


def test_instrument_export_catalogue_does_not_claim_simulated_sources(tmp_path):
    database = tmp_path / "instrument.sqlite"
    store = SessionStore(database, "instrument")
    store.append(
        Sample(
            datetime.now(UTC),
            1.0,
            "Activated",
            "instrument",
            {"LAF": 64.25},
            {"LAF": 6425},
            "unvalidated-hardware",
        )
    )
    store.close()
    target = tmp_path / "instrument.json"
    export_session(database, target, "json")
    data = json.loads(target.read_text())
    laf = next(p for p in data["parameters"] if p["id"] == "LAF")
    laeq = next(p for p in data["parameters"] if p["id"] == "LAeq")
    assert "raw / 100" in laf["conversion"]
    assert laeq["source"] == "not acquired"
    assert data["samples"][0]["raw_values"]["LAF"] == 6425

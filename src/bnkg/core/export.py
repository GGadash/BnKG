"""Stream session rows into open formats with explicit origin and parameter metadata."""

import csv
import json
import os
import sqlite3
import tempfile
import time
from contextlib import closing
from dataclasses import asdict
from datetime import UTC, datetime
from html import escape
from pathlib import Path

from openpyxl import Workbook

from bnkg.core.history import object_json
from bnkg.core.measurements import decode_outputs
from bnkg.core.parameters import PARAMETERS
from bnkg.core.spectrum import SpectrumFrame


def spreadsheet_text(value: object) -> object:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class ExportCancelled(RuntimeError):
    pass


def check_cancel(cancelled):
    if cancelled():
        raise ExportCancelled("Export cancelled; no destination was published.")


def export_session(
    database: Path,
    target: Path,
    format_name: str,
    *,
    cancelled=lambda: False,
    progress=lambda count, total, phase: None,
    parameter_ids=None,
    include_spectra=False,
    include_raw=True,
) -> int:
    """Copy a consistent SQLite snapshot, release the live DB, then format independently."""
    database, target = Path(database), Path(target)
    if database.resolve() == target.resolve():
        raise ValueError("Export must not overwrite the source database.")
    if target.exists():
        raise FileExistsError("Choose a new export name; existing files are preserved.")
    if format_name not in {"csv", "tsv", "txt", "json", "ndjson", "xlsx", "html"}:
        raise ValueError("Unsupported export format.")
    selected = [p for p in PARAMETERS if parameter_ids is None or p.id in parameter_ids]
    if not selected or (
        parameter_ids is not None and set(parameter_ids) - {p.id for p in PARAMETERS}
    ):
        raise ValueError("Choose known parameters for export.")
    if include_spectra and format_name not in {"json", "ndjson", "xlsx", "html"}:
        raise ValueError("Band data requires JSON, NDJSON, XLSX or HTML.")
    check_cancel(cancelled)
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".bnkg-export-", dir=target.parent) as directory:
        snapshot = Path(directory) / "snapshot.sqlite"
        deadline = time.monotonic() + 120

        def copying(status, remaining, total):
            check_cancel(cancelled)
            if time.monotonic() > deadline:
                raise RuntimeError("Snapshot copy timed out; try a closed session or retry later.")
            progress(total - remaining, total, "Copying database snapshot")

        with closing(
            sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True, timeout=0.25)
        ) as source:
            source.execute("PRAGMA query_only=ON")
            source.execute("PRAGMA trusted_schema=OFF")
            with closing(sqlite3.connect(snapshot)) as copy:
                source.backup(copy, pages=128, progress=copying, sleep=0.01)
        check_cancel(cancelled)
        return _export_snapshot(
            snapshot,
            target,
            format_name,
            selected,
            cancelled,
            progress,
            include_spectra,
            include_raw,
            parameter_ids is None or set(parameter_ids) == {p.id for p in PARAMETERS},
        )


def spectral_payload(raw, origin):
    context = raw.get("_hbkg_context", {})
    if (
        origin != "instrument"
        or not isinstance(context, dict)
        or not isinstance(context.get("sequences"), dict)
    ):
        return {
            "available": False,
            "reason": "No retained instrument descriptor context",
            "series": [],
        }
    sequences, setup = context["sequences"], context.get("setup_basis", {})
    if not isinstance(setup, dict) or any(not isinstance(d, dict) for d in sequences.values()):
        raise ValueError("Malformed saved spectrum context.")
    broadband, spectra = decode_outputs(raw, sequences, setup)
    frame = SpectrumFrame.capture(spectra, broadband=broadband)
    return {
        "available": bool(frame.series),
        "source": "instrument",
        "measurement_context": context.get("metadata", {}),
        "series": [
            {
                "key": series.key,
                **series.as_dict(),
                "meaning": asdict(series.meaning),
                "label": series.label,
                "broadband": broadband.get(series.meaning.quantity)
                if series.meaning.kind in {"current", "total"}
                else None,
            }
            for series in frame.series
        ],
    }


def _export_snapshot(
    database,
    target,
    format_name,
    selected,
    cancelled,
    progress,
    include_spectra,
    include_raw,
    preserve_all_values,
):
    if format_name not in {"csv", "tsv", "txt", "json", "ndjson", "xlsx", "html"}:
        raise ValueError("Unsupported export format.")
    if database.resolve() == target.resolve():
        raise ValueError("Export must not overwrite the source database.")
    if target.exists():
        raise FileExistsError("Choose a new export name; existing files are preserved.")
    target.parent.mkdir(parents=True, exist_ok=True)
    columns = ["Timestamp_UTC", "Elapsed_s", "State", "Origin", "Quality"] + [
        p.id + "_dB" for p in selected
    ]
    count = 0
    temporary: Path | None = None
    workbook = None
    workbook_saved = False
    text_chunks = None
    text_rows = text_refs = 0

    def xlsx_cell(value):
        # Excel limits UTF-16 cell length. Keep oversized text losslessly in
        # explicit, ordered JSON string chunks, never truncate provenance.
        nonlocal text_chunks, text_rows, text_refs
        if not isinstance(value, str) or len(value.encode("utf-16-le")) // 2 <= 32000:
            return spreadsheet_text(value)
        parts = (len(value) + 3999) // 4000
        if text_rows + parts > 1_048_575:
            raise ValueError("XLSX text row limit reached; use JSON/NDJSON.")
        if text_chunks is None:
            text_chunks = workbook.create_sheet("Text chunks")
            text_chunks.append(
                [
                    "Reference",
                    "Part",
                    "Parts",
                    "JSON string: decode each part, then concatenate in order",
                ]
            )
        text_refs += 1
        reference = f"text-{text_refs}"
        for index, start in enumerate(range(0, len(value), 4000), 1):
            check_cancel(cancelled)
            text_chunks.append(
                [
                    reference,
                    index,
                    parts,
                    json.dumps(value[start : start + 4000], ensure_ascii=False),
                ]
            )
        text_rows += parts
        return f"Text chunks:{reference} (JSON-decode parts, concatenate in order)"

    try:
        # Read-only URI avoids creating an empty database for a mistyped input.
        with closing(sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True)) as db:
            db.execute("PRAGMA query_only=ON")
            db.execute("PRAGMA trusted_schema=OFF")
            db.setlimit(sqlite3.SQLITE_LIMIT_LENGTH, 2_065_536)
            db.set_progress_handler(lambda: int(cancelled()), 1000)
            db.execute("BEGIN")
            if db.execute("PRAGMA user_version").fetchone()[0] != 1:
                raise ValueError("Unsupported session database version.")
            metadata = dict(db.execute("SELECT key,value FROM metadata"))
            total = db.execute("SELECT count(*) FROM samples").fetchone()[0]
            metadata["export_snapshot_utc"] = datetime.now(UTC).isoformat()
            metadata["export_scope"] = (
                "Consistent local snapshot; excludes samples appended after snapshot completion"
            )
            catalogue = [asdict(p) for p in selected]
            for parameter in catalogue:
                parameter["measurement_scope"] = parameter["description"]
                parameter["calculated"] = bool(parameter["formula"])
                if metadata["origin"] == "simulation":
                    parameter["source"] = "simulation"
                    parameter["conversion"] = "synthetic dB; no raw instrument conversion"
                elif "parameter_basis" in metadata:
                    parameter["source"] = (
                        "SLM output if present; local derived fields not populated"
                    )
                    parameter["conversion"] = metadata["parameter_basis"]
                elif parameter["id"] == "LAF":
                    parameter["conversion"] = "raw / 100; brief reference, hardware unvalidated"
                else:
                    parameter["source"] = "not acquired"
                    parameter["conversion"] = "not applicable"
            with tempfile.NamedTemporaryFile(
                dir=target.parent, suffix=".partial", delete=False
            ) as temp:
                temporary = Path(temp.name)
            sheet = None
            bands = None
            band_count = 0
            if format_name == "xlsx":
                workbook = Workbook(write_only=True)
                summary = workbook.create_sheet("Summary")
                summary.append(["Field", "Value"])
                for key, value in metadata.items():
                    summary.append([xlsx_cell(key), xlsx_cell(value)])
                parameters = workbook.create_sheet("Parameters")
                parameters.append(list(catalogue[0]))
                for parameter in catalogue:
                    parameters.append([xlsx_cell(v) for v in parameter.values()])
                sheet = workbook.create_sheet("Time Series")
                sheet.append(columns)
                if include_spectra:
                    bands = workbook.create_sheet("Spectra")
                    bands.append(
                        [
                            "Sample",
                            "Timestamp_UTC",
                            "Origin",
                            "Series",
                            "Quantity",
                            "Scope",
                            "Weighting",
                            "Integration_ms",
                            "Output_ms",
                            "Band_Hz",
                            "Level_dB",
                            "Broadband_dB",
                            "Descriptors_JSON",
                            "Measurement_context_JSON",
                            "Quality",
                            "State",
                        ]
                    )
            with temporary.open("w", encoding="utf-8", newline="") as handle:
                writer = None
                if format_name in {"csv", "tsv", "txt"}:
                    writer = csv.writer(handle, delimiter="," if format_name == "csv" else "\t")
                    writer.writerow(columns)
                elif format_name == "html":
                    handle.write(
                        '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>B&amp;KG measurement report</title><style>body{font:14px system-ui;color:#17241e;margin:2rem}table{border-collapse:collapse;max-width:100%;display:block;overflow:auto}th,td{border:1px solid #bdc8be;padding:.4rem;text-align:left}th{background:#edf3ed}code{white-space:pre-wrap}section{margin-bottom:2rem}</style><h1>B&amp;KG measurement report</h1><p>Development output. Verify against the instrument. No compliance or certification claim. PC receipt times are not integration periods.</p>'
                    )
                    handle.write(
                        "<h2>Session and parameter context</h2><pre>"
                        + escape(
                            json.dumps({"session": metadata, "parameters": catalogue}, indent=2)
                        )
                        + "</pre><h2>Recorded samples</h2>"
                    )
                elif format_name == "json":
                    handle.write(
                        '{"session":'
                        + json.dumps(metadata)
                        + ',"parameters":'
                        + json.dumps(catalogue)
                        + ',"samples":['
                    )
                elif format_name == "ndjson":
                    handle.write(
                        json.dumps(
                            {"type": "metadata", "session": metadata, "parameters": catalogue}
                        )
                        + "\n"
                    )
                for ident, timestamp, elapsed, state, origin, quality, encoded, raw in db.execute(
                    "SELECT id,timestamp,elapsed_s,state,origin,quality,values_json,raw_json FROM samples ORDER BY id"
                ):
                    check_cancel(cancelled)
                    values = object_json(encoded, 65536)
                    raw_values = object_json(raw)
                    if not preserve_all_values:
                        values = {p.id: values.get(p.id) for p in selected}
                    spectra = spectral_payload(raw_values, origin) if include_spectra else None
                    row = [timestamp, elapsed, state, origin, quality] + [
                        values.get(p.id) for p in selected
                    ]
                    sample = {
                        "id": ident,
                        "timestamp": timestamp,
                        "elapsed_s": elapsed,
                        "state": state,
                        "origin": origin,
                        "quality": quality,
                        "values": values,
                        **({"raw_values": raw_values} if include_raw else {}),
                        **({"spectra": spectra} if include_spectra else {}),
                    }
                    if bands is not None:
                        context_cell = xlsx_cell(json.dumps(spectra.get("measurement_context", {})))
                        for series in spectra["series"]:
                            descriptor_cell = xlsx_cell(json.dumps(series["descriptor"]))
                            meaning = series["meaning"]
                            for frequency, value in zip(
                                series["frequencies"], series["values"], strict=True
                            ):
                                check_cancel(cancelled)
                                band_count += 1
                                if band_count > 1_048_575:
                                    raise ValueError(
                                        "XLSX spectrum row limit reached; use JSON or NDJSON."
                                    )
                                cells = [
                                    ident,
                                    timestamp,
                                    origin,
                                    series["key"],
                                    meaning["quantity"],
                                    meaning["kind"],
                                    meaning["weighting"],
                                    meaning["period_ms"],
                                    meaning["output_ms"],
                                    frequency,
                                    value,
                                    series["broadband"],
                                    descriptor_cell,
                                    context_cell,
                                    quality,
                                    state,
                                ]
                                bands.append([xlsx_cell(v) for v in cells])
                    if writer:
                        writer.writerow([spreadsheet_text(value) for value in row])
                    elif sheet is not None:
                        if count >= 1_048_575:
                            raise ValueError(
                                "XLSX row limit reached; use NDJSON or CSV for this session."
                            )
                        sheet.append([xlsx_cell(value) for value in row])
                    elif format_name == "html":
                        handle.write(
                            "<section><h3>Sample "
                            + str(ident)
                            + "</h3><p>"
                            + escape(
                                str(timestamp)
                                + " · "
                                + str(state)
                                + " · "
                                + str(origin)
                                + " · "
                                + str(quality)
                            )
                            + "</p><table><tr><th>Parameter</th><th>Level dB</th><th>Meaning</th></tr>"
                        )
                        for p in selected:
                            value = values.get(p.id)
                            handle.write(
                                "<tr><td>"
                                + escape(p.label + (" · Calculated" if p.formula else ""))
                                + "</td><td>"
                                + escape("Unavailable" if value is None else str(value))
                                + "</td><td>"
                                + escape(p.description)
                                + "</td></tr>"
                            )
                        handle.write("</table>")
                        if include_spectra:
                            handle.write(
                                "<details><summary>Direct saved spectra and descriptors</summary><pre>"
                                + escape(json.dumps(spectra, indent=2))
                                + "</pre></details>"
                            )
                        handle.write("</section>")
                    elif format_name == "json":
                        handle.write(("," if count else "") + json.dumps(sample, allow_nan=False))
                    elif format_name == "ndjson":
                        handle.write(
                            json.dumps({"type": "sample", **sample}, allow_nan=False) + "\n"
                        )
                    count += 1
                    if count % 100 == 0 or count == total:
                        progress(count, total, "Writing " + format_name.upper())
                if format_name == "json":
                    handle.write("]}")
                if format_name == "html":
                    handle.write("</html>")
            check_cancel(cancelled)
            if workbook:
                events = workbook.create_sheet("Events")
                events.append(["Timestamp_UTC", "Kind", "Detail"])
                for row in db.execute("SELECT timestamp,kind,detail FROM events"):
                    check_cancel(cancelled)
                    events.append([xlsx_cell(v) for v in row])
                progress(count, total, "Packaging XLSX")
                workbook.save(temporary)
                workbook_saved = True
        # Atomic hard-link publication refuses existing targets on Windows and Linux.
        if target.exists():
            raise FileExistsError("Destination appeared during export; choose another filename.")
        check_cancel(cancelled)
        os.link(temporary, target)
        return count
    except sqlite3.Error:
        check_cancel(cancelled)
        raise
    finally:
        if workbook is not None:
            if not workbook_saved and temporary:
                # Complete write-only streams so openpyxl cleans its temporary XML files.
                # Cancellation during ZIP finalization is honored before publication.
                try:
                    workbook.save(temporary)
                except Exception:
                    pass
            workbook.close()
        if temporary and temporary.exists():
            temporary.unlink()
